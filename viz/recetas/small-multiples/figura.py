"""
Receta «small multiples» — caso de referencia: PIB sectorial del primer semestre.

Doce paneles en un solo lienzo (4 x 3): el PIB primero, como ancla, y los once
sectores del INE ordenados de mas a menos recuperados en 2026. Cada panel lleva
SU PROPIA escala: la pregunta de esta lamina es el perfil de cada sector, no la
comparacion de niveles entre paneles.

Produce DOS laminas del mismo molde:

  indice     nivel del semestre I+II con 2019 = 100 (base prepandemia)
  variacion  variacion interanual del semestre I+II, en porcentaje

El año 2026 no es dato de cuentas nacionales: es el semestre de 2025 movido con
las tasas del IpAEC del BCB. Lo marca el AREA SOMBREADA, y solo eso: la linea se
dibuja continua.

Ejecutar:  python figura.py            -> output/*.png  (las dos laminas)
Publicar:  python publicar.py          (despues de este)
"""
import re
import sys
import unicodedata
from pathlib import Path

import numpy as np
import openpyxl

PUB = Path(__file__).resolve().parent
VIZ = PUB.parents[1]                       # .../galeria-populi/viz
sys.path.insert(0, str(VIZ))
import populi_style as ps                  # noqa: E402

FORMATO = "informe_mosaico"
XLSX = PUB / "data" / "populi_pib_ipaec_2026_4.xlsx"
SHEET = "Sectorial semestre"
OUT = PUB / "output"

COLS, ROWS = 4, 3

# Canal LIBRE de verdad entre columna y columna, en px @1080 (se escala por sc,
# como todo el motor). Es lo UNICO que se fija a mano del reparto horizontal: el
# canal de cifras y el desborde de la etiqueta de año se MIDEN, y el ancho de
# trazado se despeja de ahi.
CANAL = 46
EST_YEAR = 2026               # primer (y unico) año estimado con el IpAEC

# Codificacion por SIGNO del AREA, no de la linea. La linea es una sola, en
# tinta: lo que se colorea es la BRECHA contra la base, que es justamente lo que
# esta lamina viene a mostrar. (La regla de "un solo color" del Banco se fijo
# para BARRAS, donde el color era la marca misma.)
C_LINEA = ps.col("tinta")
C_POS = ps.col("serie_teal")
C_NEG = ps.col("rojo")
A_POS, A_NEG = 0.22, 0.18

# Escala tipografica propia de la figura. SIZES esta calibrado para UN panel por
# lienzo; aqui hay DOCE. Muy por debajo de la global del motor.
S_EJE = 15        # cifras de los ejes             (motor: 25)
S_ROTULO = 19     # rotulo de cada panel           (motor: 29)
S_DATO = 15       # cifra del ultimo año           (motor: 23)

# Nombre largo del INE -> rotulo del panel, en Cada Palabra En Mayuscula.
CORTO = {
    "Actividad extractiva": "Actividad Extractiva",
    "Agricultura, ganadería, silvicultura y pesca": "Agropecuaria y Pesca",
    "Electricidad, agua y recolección de desechos": "Electricidad y Agua",
    "Administración pública, salud y educación": "Administración Pública",
    "Actividades comunales, sociales y personales": "Comunales y Personales",
    "Alojamiento y servicio de comidas y bebidas": "Alojamiento y Comidas",
    "Financieras, seguros, inmobiliarias y prof.": "Financieras e Inmobiliarias",
    "Industrias manufactureras": "Industria Manufacturera",
    "Comercio": "Comercio",
    "Transporte y comunicaciones": "Transporte y Comunicaciones",
    "Construcción": "Construcción",
    "PRODUCTO INTERNO BRUTO": "Producto Interno Bruto",
}
PIB_KEY = "PRODUCTO INTERNO BRUTO"


# --------------------------------------------------------------------------- #
# Datos
# --------------------------------------------------------------------------- #
def _norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c)).lower().strip()


def _bloque(ws, titulo):
    """Fila del encabezado de tabla que sigue al titulo del bloque.

    Se busca por TEXTO y no por numero de fila: el libro se edita a mano y las
    filas se corren de una version a otra.
    """
    key = _norm(titulo)
    for r in range(1, ws.max_row + 1):
        if key in _norm(ws.cell(r, 1).value):
            return r + 1
    raise ValueError("no encuentro el bloque «%s» en la hoja %s" % (titulo, SHEET))


def _tabla(ws, fila, escala=1.0):
    """(años, {sector: [valores]}) leyendo hasta la primera fila vacia."""
    años, cols = [], []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(fila, c).value
        # Los encabezados de año estan como TEXTO en el libro, no como numero.
        if v is not None and re.fullmatch(r"(19|20)\d{2}", str(v).strip()):
            años.append(int(str(v).strip()))
            cols.append(c)
    if not años:
        raise ValueError("sin columnas de año en la fila %d" % fila)

    serie, r = {}, fila + 1
    while r <= ws.max_row:
        nombre = ws.cell(r, 1).value
        if nombre is None or not str(nombre).strip():
            break
        vals = [ws.cell(r, c).value for c in cols]
        if all(isinstance(v, (int, float)) for v in vals):
            serie[str(nombre).strip()] = [float(v) * escala for v in vals]
        r += 1
    return años, serie


def leer():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[SHEET]
    a_idx, s_idx = _tabla(ws, _bloque(ws, "4 · Índice del semestre"))
    # El bloque 5 guarda fracciones (0,048), no puntos porcentuales.
    a_var, s_var = _tabla(ws, _bloque(ws, "5 · Variación interanual"), escala=100.0)

    faltan = [k for k in CORTO if k not in s_idx or k not in s_var]
    if faltan:
        raise ValueError("el libro no trae: " + " | ".join(faltan))

    # PIB primero (es el ancla) y los sectores de mas a menos recuperados en 2026.
    # El MISMO orden en las dos laminas: el lector pasa de una a otra sin tener
    # que reubicar los paneles.
    orden = [PIB_KEY] + sorted((k for k in CORTO if k != PIB_KEY),
                               key=lambda k: s_idx[k][-1], reverse=True)
    return orden, (a_idx, s_idx), (a_var, s_var)


# --------------------------------------------------------------------------- #
# Escala redondeada por panel
# --------------------------------------------------------------------------- #
def rango(vals, base, pad=0.12, divs=3):
    """Min/max/paso «lindos» que contienen a los datos y a la linea base.

    Se prueban varios pasos y gana el que menos rango DESPERDICIA: quedarse con
    el primero que redondea bonito abria −75 para un minimo de −47.
    """
    lo, hi = min(min(vals), base), max(max(vals), base)
    span = (hi - lo) or 1.0
    lo, hi = lo - span * pad, hi + span * pad
    crudo = (hi - lo) / divs
    e = int(np.floor(np.log10(crudo)))
    mejor = None
    for k in (e - 1, e, e + 1):
        for m in (1, 2, 2.5, 5):
            paso = m * 10.0 ** k
            mn, mx = np.floor(lo / paso) * paso, np.ceil(hi / paso) * paso
            n = int(round((mx - mn) / paso))
            if not 3 <= n <= 6:
                continue
            if mejor is None or (mx - mn) < (mejor[1] - mejor[0]):
                mejor = (mn, mx, paso)
    return mejor or (lo, hi, (hi - lo) / divs)


# --------------------------------------------------------------------------- #
# Una lamina
# --------------------------------------------------------------------------- #
def lamina(nombre, orden, años, serie, base, sufijo, decimales,
           titulo, subtitulo, nota, ticks_x):
    W, H, sc = ps._spec(FORMATO)
    M = ps.MARGIN * sc

    fig, axd = ps.nueva_figura(FORMATO)
    axd.set_xticks([]); axd.set_yticks([])
    ps.componer(
        fig, axd,
        titulo=titulo, subtitulo=subtitulo,
        fuente="Fuente: Instituto Nacional de Estadística (INE), PIB trimestral base 2017; "
               "Banco Central de Bolivia (BCB), Reporte de Inflación y Política Monetaria "
               "(IpAEC). Elaboración: Centro de Estudios POPULI · Carlos Aranda.",
        nota=nota, formato=FORMATO,
    )
    pos = axd.get_position()
    axd.set_visible(False)

    x = np.asarray(años, dtype=float)

    # Reparto de la cuadricula. Los huecos van en px sobre el lienzo real: el
    # horizontal lo fija la etiqueta de año, que va CENTRADA sobre su marca y
    # cuelga media palabra fuera del panel por cada lado.
    gap_y = 58 * sc
    rot_h = 46 * sc                     # banda del rotulo, encima de cada panel
    grid_w, grid_h = pos.width * W, pos.height * H
    cellW = grid_w / COLS - CANAL * sc     # provisional: se recalcula al medir
    cellH = (grid_h - (ROWS - 1) * gap_y) / ROWS
    panelH = cellH - rot_h

    fp_eje = ps.fp(ps.MONO, S_EJE * sc)
    fp_rot = ps.fp("Inter Bold", S_ROTULO * sc)
    fp_dato = ps.fp("JetBrains Mono SemiBold", S_DATO * sc)

    ejes, celdas = [], []
    for i, clave in enumerate(orden):
        c, r = i % COLS, i // COLS
        cx = pos.x0 * W + c * (grid_w / COLS)
        cy = pos.y0 * H + (ROWS - 1 - r) * (cellH + gap_y)
        ax = fig.add_axes([cx / W, cy / H, cellW / W, panelH / H])
        ejes.append(ax); celdas.append((cx, cy))

        y = np.asarray(serie[clave], dtype=float)

        # El tramo estimado: area sombreada y NADA MAS. La serie va continua.
        # Medio paso antes de la primera observacion estimada y hasta el borde
        # del panel, como en las otras recetas: la banda tiene que ENCERRAR al
        # punto de 2026, no cortarse encima de el.
        ax.axvspan(EST_YEAR - 0.5, x[-1] + 0.35, color=ps.COLORS["cafe"],
                   alpha=0.085, linewidth=0, zorder=1)

        # Brecha contra la base, coloreada por signo. interpolate=True corta
        # exactamente en el cruce; sin el, el relleno inventa un escalon plano
        # entre dos años cuando la serie atraviesa la base.
        ax.fill_between(x, y, base, where=(y >= base), interpolate=True,
                        color=C_POS, alpha=A_POS, linewidth=0, zorder=2)
        ax.fill_between(x, y, base, where=(y <= base), interpolate=True,
                        color=C_NEG, alpha=A_NEG, linewidth=0, zorder=2)

        ax.axhline(base, color=ps.COLORS["tinta"], linewidth=1.15 * sc, zorder=4)

        ax.plot(x, y, color=C_LINEA, linewidth=1.7 * sc, zorder=6,
                solid_capstyle="round", solid_joinstyle="round")
        ax.plot(x, y, linestyle="none", marker="o", markersize=3.4 * sc,
                markerfacecolor=ps.COLORS["fondo"], markeredgecolor=C_LINEA,
                markeredgewidth=1.25 * sc, zorder=7)

        mn, mx, paso = rango(y, base)
        ax.set_ylim(mn, mx)
        ax.set_yticks(np.arange(mn, mx + paso / 2, paso))
        y_dec = 0 if abs(round(paso) - paso) < 1e-9 else 1
        ax.yaxis.set_major_formatter(ps.formateador_es(y_dec, ""))
        ax.set_xlim(x[0] - 0.35, x[-1] + 0.35)
        ax.set_xticks(ticks_x)
        ax.xaxis.set_major_formatter(ps.formateador_es(0, ""))

        ps.aplicar_estilo_ejes(ax, grid_y=False)
        ax.grid(axis="y", color=ps.COLORS["borde"], linewidth=1.0 * sc,
                linestyle=(0, (1.6, 2.6)), zorder=0)
        ax.set_axisbelow(True)
        for lbl in ax.get_xticklabels() + ax.get_yticklabels():
            lbl.set_fontproperties(fp_eje)

        # Cifra del ultimo año, DENTRO del panel y pegada al borde derecho, en
        # la esquina que la serie deja libre. Afuera se leia igual de bien pero
        # obligaba a reservar ~175 px por columna —el 16 % del ancho util para
        # doce numeros—; adentro ese ancho se lo quedan los paneles. La esquina
        # se elige mirando SOLO el tercio derecho, que es donde va a caer.
        der = x >= x[-1] - 0.34 * (x[-1] - x[0])
        arriba = (mx - y[der].max()) >= (y[der].min() - mn)
        ax.text(0.985, 0.94 if arriba else 0.06,
                ps.es_num(y[-1], decimales) + sufijo,
                transform=ax.transAxes, ha="right",
                va="top" if arriba else "bottom",
                fontproperties=fp_dato, color=ps.COLORS["cafe_oscuro"], zorder=8)

    # ── Encuadre: SEGUNDA PASADA, con las medidas reales ──────────────────── #
    # Los doce paneles tienen que medir EXACTAMENTE lo mismo: si cada uno se
    # ajusta a lo ancho de sus propias cifras, las columnas salen distintas y los
    # perfiles dejan de ser comparables, que es todo el punto de la lamina.
    #
    # Con un hueco fijo entre columnas sobraba aire: el canal medido daba 258 px
    # (tres veces = 17 % del ancho util) porque al hueco se le sumaba lo que cada
    # panel NO usaba de su reserva. Asi que el hueco no se fija: se fija el CANAL
    # LIBRE que se quiere ver, se mide cuanto desborda de verdad cada panel por
    # cada lado, y se despeja el ancho de trazado que cabe.
    fig.canvas.draw()
    rend = fig.canvas.get_renderer()

    izq, der = [], []                      # tinta fuera del area de trazado
    for ax in ejes:
        p = ax.get_position()
        sx, aw0 = p.x0 * W, p.width * W
        x0s = [l.get_window_extent(rend).x0 for l in ax.get_yticklabels() if l.get_text()]
        izq.append(sx - min(x0s) if x0s else 0.0)
        x1s = [l.get_window_extent(rend).x1 for l in ax.get_xticklabels() if l.get_text()]
        der.append(max(x1s) - (sx + aw0) if x1s else 0.0)

    B = [max(izq[c::COLS]) for c in range(COLS)]     # desborde izquierdo por columna
    R = [max(der[c::COLS]) for c in range(COLS)]     # desborde derecho por columna

    # paso = ancho + K; K sale del canal libre exigido entre columna y columna.
    K = max(R[c] + B[c + 1] for c in range(COLS - 1)) + CANAL * sc
    ancho_util = (grid_w - B[0] - R[-1] - (COLS - 1) * K) / COLS
    paso = ancho_util + K

    columnas = [pos.x0 * W + B[0] + c * paso for c in range(COLS)]   # x de cada eje
    celdas = [(columnas[i % COLS] - B[i % COLS], celdas[i][1])       # x de cada rotulo
              for i in range(len(ejes))]
    for i, ax in enumerate(ejes):
        p = ax.get_position()
        ax.set_position([columnas[i % COLS] / W, p.y0, ancho_util / W, p.height])

    # Rotulo de cada panel, al borde de su celda, sobre el area de trazado. Al
    # borde de la CELDA y no del area de trazado: asi el del primer panel cae en
    # el margen del titulo y los tres de cada columna quedan a plomo.
    for (cx, cy), clave in zip(celdas, orden):
        fig.text(cx / W, (cy + panelH + 13 * sc) / H, CORTO[clave],
                 fontproperties=fp_rot, color=ps.COLORS["cafe_oscuro"],
                 va="bottom", ha="left")

    destino = OUT / f"pib_semestre_{nombre}.png"
    ps.guardar(fig, str(destino), formato=FORMATO, thumb=False)
    return destino


# --------------------------------------------------------------------------- #
def main():
    orden, (a_idx, s_idx), (a_var, s_var) = leer()
    print("%d paneles | indice %d-%d | variacion %d-%d"
          % (len(orden), a_idx[0], a_idx[-1], a_var[0], a_var[-1]))

    lamina(
        "indice_2019", orden, a_idx, s_idx, base=100.0, sufijo="", decimales=1,
        titulo="Bolivia: Nivel del PIB Sectorial Frente a la Prepandemia",
        subtitulo="Índice del primer semestre (enero–junio), 2019 = 100",
        nota="Área turquesa: por encima del nivel de 2019; roja: por debajo. Cada panel "
             "tiene su propia escala. El semestre de 2026 (área sombreada) es estimación "
             "a partir del IpAEC del BCB.",
        ticks_x=[y for y in a_idx if (y - a_idx[0]) % 3 == 0],
    )
    lamina(
        "variacion", orden, a_var, s_var, base=0.0, sufijo="%", decimales=1,
        titulo="Bolivia: Crecimiento del PIB Sectorial",
        subtitulo="Variación interanual del primer semestre (enero–junio), en porcentaje",
        nota="Área turquesa: crecimiento; roja: caída. Cada panel tiene su propia escala. "
             "El semestre de 2026 (área sombreada) es estimación a partir del IpAEC del BCB.",
        ticks_x=[y for y in a_var if (y - a_var[0]) % 2 == 0],
    )


if __name__ == "__main__":
    main()
