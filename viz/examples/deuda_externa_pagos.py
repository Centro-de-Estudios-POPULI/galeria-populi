"""
Lámina del Banco: pagos de la deuda externa bruta de Bolivia por trimestre, capital e interés.

    python viz/examples/deuda_externa_pagos.py                 # la lámina + manifiesto + contrato
    python viz/examples/deuda_externa_pagos.py --sin-contrato  # iterar el diseño

Datos: el Excel que armó Carlos con datos del Banco Central de Bolivia («pagos de deuda.xlsx», en el
Escritorio; millones de dólares). Se lee una COPIA, porque Excel bloquea el archivo mientras está abierto,
y antes de dibujar se comprueba que capital + interés = total en cada trimestre: si no cuadra, no hay
lámina. Las cifras de la nota de la ficha salen de las mismas filas. Nada de datos al repo: el Banco
publica la imagen. Tipografía del blog (Playfair Display 700 + Inter), como `blog_tarifas.py`.
"""
import argparse
import itertools
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
from matplotlib.lines import Line2D
from matplotlib.patches import Rectangle
from PIL import Image

VIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(VIZ))
import populi_style as ps   # noqa: E402
import catalogo as CAT      # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

EXCEL = Path.home() / "OneDrive" / "Desktop" / "pagos de deuda.xlsx"
FECHA = "2026-10-02"
FUENTE = "Fuente: Banco Central de Bolivia (BCB). Elaboración: Centro de Estudios POPULI · Carlos Aranda."
SLUG = "deuda-externa-bruta-pagos-trimestrales"

TIT = "Playfair Display Bold"          # el titular del blog
ps.set_tema("Inter", "Inter")          # cifras en Inter, como en el blog
P = ps._pal
ROJO, TURQ = P.BRAND, P.FAMILIES["turquesa"][1]          # el par validado: #C71E1D y #0A9396
TINTA, PIZARRA, GRILLA = P.INK, P.MUTED, P.GRID
FONDO = ps.COLORS["fondo"]


# --------------------------------------------------------------------------- #
# Datos
# --------------------------------------------------------------------------- #
def leer():
    """Las tres filas del Excel (total, capital, interés) con su trimestre y año."""
    with tempfile.TemporaryDirectory() as d:
        copia = Path(d) / "pagos.xlsx"
        shutil.copy2(EXCEL, copia)
        ws = openpyxl.load_workbook(copia, data_only=True).active
        filas = {str(r[3].value).strip(): r for r in ws.iter_rows() if r[3].value}
        anios = next(r for r in ws.iter_rows() if any(isinstance(c.value, int) and 2000 < c.value < 2100 for c in r))
        trims = ws[anios[0].row + 1]
        anio, cols = None, []
        for c in anios[4:]:
            anio = c.value if c.value else anio     # celdas combinadas: el año vale para las que siguen
            if trims[c.column - 1].value:
                cols.append((c.column - 1, anio, str(trims[c.column - 1].value).strip()))
        dato = lambda nombre: [float(filas[nombre][i].value) for i, _, _ in cols]
        tot, cap, inte = dato("Pago total de Deuda Externa Bruta"), dato("Capital"), dato("Interés")
    for t, c, i, (_, a, q) in zip(tot, cap, inte, cols):
        assert abs(c + i - t) < 1e-6, f"{q} {a}: capital + interés ≠ total ({c} + {i} ≠ {t})"
    return [{"anio": a, "trim": q, "total": t, "capital": c, "interes": i}
            for (_, a, q), t, c, i in zip(cols, tot, cap, inte)]


# --------------------------------------------------------------------------- #
# Piezas del encuadre (las mismas reglas que blog_tarifas.py)
# --------------------------------------------------------------------------- #
def choques(fig, nombre, W, margen):
    """Ningún texto se pisa con otro ni sale del margen lateral."""
    fig.canvas.draw()
    rr = fig.canvas.get_renderer()
    textos = list(fig.texts) + [t for lg in fig.legends for t in lg.get_texts()]
    for a in fig.axes:
        textos += list(a.texts)
        if a.axison:
            (lx, hx), (ly, hy) = sorted(a.get_xlim()), sorted(a.get_ylim())
            textos += [t for t in a.get_xticklabels() if lx <= t.get_position()[0] <= hx]
            textos += [t for t in a.get_yticklabels() if ly <= t.get_position()[1] <= hy]
    cajas = [(t.get_text().replace("\n", " ")[:40], t.get_window_extent(rr))
             for t in textos if t.get_visible() and t.get_text().strip()]
    malos = [f"fuera del margen: «{x}»" for x, b in cajas if b.x0 < margen - 1 or b.x1 > W - margen + 1]
    for (t1, b1), (t2, b2) in itertools.combinations(cajas, 2):
        if b1.x0 < b2.x1 - 1 and b2.x0 < b1.x1 - 1 and b1.y0 < b2.y1 - 1 and b2.y0 < b1.y1 - 1:
            malos.append(f"se pisan: «{t1}» y «{t2}»")
    print(f"  {nombre}: {len(cajas)} textos, " + ("sin choques" if not malos else f"{len(malos)} PROBLEMAS"))
    for m in malos:
        print("    ⚠ " + m)
    return not malos


def tinta_fuera(png, margen):
    """Sobre el PNG, no a ojo: ninguna tinta fuera de [M, W − M] (la regla del wordmark)."""
    im = np.asarray(Image.open(png).convert("RGB")).astype(int)
    fondo = np.array([int(FONDO[i:i + 2], 16) for i in (1, 3, 5)])
    tinta = (np.abs(im - fondo).sum(axis=2) > 24).any(axis=0)
    m = int(round(margen))
    fuera = int(tinta[:m - 1].sum() + tinta[im.shape[1] - m + 1:].sum())
    print(f"  tinta fuera del margen: {'ninguna' if not fuera else f'⚠ {fuera} columnas'}")
    return not fuera


def encajar_eje(fig, ax, limite_px):
    """El eje termina donde termina el wordmark (W − M): se mide el último rótulo y la línea base."""
    W = fig.bbox.width
    for _ in range(6):
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        ext = [ax.spines["bottom"].get_window_extent(r).x1]
        ext += [l.get_window_extent(r).x1 for l in ax.get_xticklabels() if l.get_text()]
        over = max(ext) - limite_px
        if over <= 0.5:
            return
        p = ax.get_position()
        ax.set_position([p.x0, p.y0, p.width - (over + 1) / W, p.height])


def leyenda_arriba(fig, ax, series, W, H, sc, M):
    """Leyenda sobre el gráfico, alineada al margen; el eje le cede ese alto."""
    p = ax.get_position()
    s_ley = ps.SIZES["leyenda"] * sc
    hs = [Rectangle((0, 0), 1, 1, facecolor=c, edgecolor="none") for _, c in series]
    fig.legend(hs, [n for n, _ in series], loc="upper left", bbox_to_anchor=(M / W, (p.y1 * H + 8 * sc) / H),
               ncol=len(series), frameon=False, prop=ps.fp(ps.BODY, s_ley), labelcolor=TINTA,
               handletextpad=0.55, columnspacing=1.6, borderaxespad=0, borderpad=0, handlelength=1.0,
               handleheight=1.0)
    ax.set_position([p.x0, p.y0, p.width, p.height - (s_ley * 1.55 + 34 * sc) / H])


# --------------------------------------------------------------------------- #
# La lámina
# --------------------------------------------------------------------------- #
def lamina(datos):
    F = "informe_5x4"
    W, H, sc = ps._spec(F)
    M = ps.MARGIN * sc
    n = len(datos)
    x = np.arange(n)
    cap = np.array([d["capital"] for d in datos])
    inte = np.array([d["interes"] for d in datos])
    tot = np.array([d["total"] for d in datos])
    ancho_barra = 0.62
    s_dato = ps.SIZES["dato"] * sc

    fig, ax = ps.nueva_figura(F)
    ax.bar(x, cap, width=ancho_barra, color=ROJO, linewidth=0, zorder=3)
    ax.bar(x, inte, width=ancho_barra, bottom=cap, color=TURQ, linewidth=0, zorder=3)
    tope = float(np.ceil(tot.max() * 1.1 / 10) * 10)   # el aire justo para el total más alto
    for xi, c, i, t in zip(x, cap, inte, tot):
        # cifras dentro de cada tramo (blanco sobre el color) y el total arriba, en negrita
        for y0, alto in ((0, c), (c, i)):
            if alto / tope * 100 > 4.2:
                ax.text(xi, y0 + alto / 2, ps.es_num(alto, 1, miles=True), ha="center", va="center", color="#FFFFFF",
                        fontproperties=ps.fp(ps.BODY, s_dato * 0.92), zorder=4)
        ax.text(xi, t + tope * 0.012, ps.es_num(t, 1, miles=True), ha="center", va="bottom", color=TINTA,
                fontproperties=ps.fp(ps.BOLD, s_dato * 1.04), zorder=4)

    ps.aplicar_estilo_ejes(ax, grid_y=False)
    ax.spines["bottom"].set_capstyle("butt")
    ax.grid(axis="y", color=GRILLA, linewidth=1.0 * sc, linestyle=(0, (1.6, 2.6)), zorder=0)
    ax.set_axisbelow(True)
    ax.set_xlim(-0.6, n - 0.4)
    ax.set_ylim(0, tope)
    ax.set_yticks(np.arange(0, tope, 200))
    ax.yaxis.set_major_formatter(ps.formateador_es(0, miles=True))
    ax.set_xticks(x)
    ax.set_xticklabels([d["trim"] for d in datos])

    sem26 = sum(d["total"] for d in datos if d["anio"] == 2026)
    a27 = sum(d["total"] for d in datos if d["anio"] == 2027)
    sem28 = sum(d["total"] for d in datos if d["anio"] == 2028)
    p_cap = cap.sum() / tot.sum() * 100
    pico = max(datos, key=lambda d: d["total"])
    meta = {"slug": SLUG, "titulo": "Bolivia: Pagos de la deuda externa bruta",
            "subtitulo": "Capital e interés, por trimestre · Millones de dólares",
            "categoria": "fiscal", "fuente": FUENTE, "fecha": FECHA, "datos": None, "grupo": "Deuda externa",
            "tags": ["deuda externa", "servicio de la deuda", "capital", "intereses", "BCB", "Bolivia"],
            "nota": (f"Pago de capital e interés de la deuda externa bruta, por trimestre calendario, de "
                     f"{datos[0]['trim']} de {datos[0]['anio']} a {datos[-1]['trim']} de {datos[-1]['anio']}. "
                     f"Suman {ps.es_num(sem26, 1, miles=True)} millones de dólares en el segundo semestre de 2026, "
                     f"{ps.es_num(a27, 1, miles=True)} millones en 2027 y {ps.es_num(sem28, 1, miles=True)} millones en el primer "
                     f"semestre de 2028. El capital es el {ps.es_num(p_cap, 0)} % de lo que se paga en todo el "
                     f"período; el trimestre más pesado es {pico['trim']} de {pico['anio']}, con "
                     f"{ps.es_num(pico['total'], 1, miles=True)} millones."),
            "tipo": "barras_apiladas", "formato": F}

    ps.componer(fig, ax, titulo=meta["titulo"], subtitulo=meta["subtitulo"], fuente=FUENTE, nota="",
                formato=F, titulo_familia=TIT)
    leyenda_arriba(fig, ax, [("Capital", ROJO), ("Interés", TURQ)], W, H, sc, M)
    # el renglón de los años, debajo de los trimestres. El motor ya deja lugar para UN renglón de rótulos;
    # el de los años se gana subiendo el eje lo justo: se mide contra el filete rojo del wordmark (el
    # elemento más alto del pie) y se deja 26 px de aire, ni más ni menos.
    s_eje = ps.SIZES["eje"] * sc
    filete_rojo = max(a.get_ydata()[0] for a in fig.artists if isinstance(a, Line2D)) * H
    anios = []
    for _ in range(4):
        encajar_eje(fig, ax, W - M)
        for a in anios:
            a.remove()
        anios = []
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        piso = min(l.get_window_extent(r).y0 for l in ax.get_xticklabels())
        y_filete = piso - 12 * sc
        for anio, grupo in itertools.groupby(enumerate(datos), key=lambda e: e[1]["anio"]):
            idx = [i for i, _ in grupo]
            x0 = ax.transData.transform((idx[0] - ancho_barra / 2, 0))[0]
            x1 = ax.transData.transform((idx[-1] + ancho_barra / 2, 0))[0]
            anios.append(fig.add_artist(Line2D([x0 / W, x1 / W], [y_filete / H, y_filete / H],
                                               transform=fig.transFigure, color=PIZARRA, linewidth=1.2 * sc,
                                               solid_capstyle="butt")))
            anios.append(fig.text((x0 + x1) / 2 / W, (y_filete - 7 * sc) / H, str(anio), ha="center", va="top",
                                  color=TINTA, fontproperties=ps.fp(ps.BOLD, s_eje)))
        fig.canvas.draw()
        suelo = min(t.get_window_extent(fig.canvas.get_renderer()).y0 for t in anios if hasattr(t, "get_text"))
        falta = (filete_rojo + 26 * sc) - suelo           # > 0: los años bajan demasiado; < 0: sobra aire
        if abs(falta) <= 1:
            break
        p = ax.get_position()
        ax.set_position([p.x0, p.y0 + falta / H, p.width, p.height - falta / H])

    ok = choques(fig, SLUG, W, M)
    png = CAT.GRAFICAS / f"{SLUG}.png"
    ps.guardar(fig, png, formato=F)
    ok = tinta_fuera(png, M) and ok
    CAT.registrar(meta, meta["tipo"], F)
    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sin-contrato", action="store_true", help="no corre verificar.py (iterar el diseño)")
    args = ap.parse_args()
    datos = leer()
    print(f"Datos: {len(datos)} trimestres, capital + interés = total en todos")
    if not lamina(datos):
        raise SystemExit("⛔ la lámina tiene textos encimados o tinta fuera del margen")
    CAT.build_manifest()
    if args.sin_contrato:
        return
    r = subprocess.run([sys.executable, str(VIZ / "verificar.py")], cwd=VIZ.parent,
                       capture_output=True, text=True, encoding="utf-8")
    print("\n".join(r.stdout.strip().splitlines()[-3:]))
    if r.returncode:
        raise SystemExit("⛔ el contrato del Banco no pasa:\n" + r.stdout[-1500:] + r.stderr[-800:])


if __name__ == "__main__":
    main()
